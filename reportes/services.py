"""
Servicios de exportación de reportes (PDF y Excel).
"""
from io import BytesIO
from decimal import Decimal
from django.http import HttpResponse
from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa


def generar_excel_financiero(kpis, fin_data, fecha_inicio, fecha_fin):
    """
    Genera archivo Excel con reporte financiero ejecutivo.
    
    Args:
        kpis (dict): KPIs globales
        fin_data (dict): datos financieros por película
        fecha_inicio (str): fecha de inicio en formato ISO
        fecha_fin (str): fecha de fin en formato ISO
    
    Returns:
        BytesIO: buffer con el archivo Excel
    """
    wb = Workbook()
    
    # Hoja 1: Resumen Global
    ws1 = wb.active
    ws1.title = "Resumen Global"
    
    # Header styling
    header_fill = PatternFill(start_color="a855f7", end_color="a855f7", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    
    ws1['A1'] = "Reporte Financiero Ejecutivo"
    ws1['A1'].font = Font(bold=True, size=14)
    ws1.merge_cells('A1:C1')
    
    ws1['A2'] = f"Período: {fecha_inicio} a {fecha_fin}"
    ws1.merge_cells('A2:C2')
    
    ws1['A4'] = "Métrica"
    ws1['B4'] = "Valor"
    for cell in ['A4', 'B4']:
        ws1[cell].fill = header_fill
        ws1[cell].font = header_font
        ws1[cell].alignment = Alignment(horizontal='center')
    
    ws1['A5'] = "Total Recaudación"
    ws1['B5'] = kpis['total_recaudacion']
    ws1['B5'].number_format = '$#,##0.00'
    
    ws1['A6'] = "Total Tickets"
    ws1['B6'] = kpis['total_tickets']
    ws1['B6'].number_format = '#,##0'
    
    ws1['A7'] = "Ticket Promedio"
    ws1['B7'] = kpis['ticket_promedio']
    ws1['B7'].number_format = '$#,##0.00'
    
    # Ajustar anchos de columna
    ws1.column_dimensions['A'].width = 25
    ws1.column_dimensions['B'].width = 18
    
    # Hoja 2: Por Película
    ws2 = wb.create_sheet("Por Película")
    ws2['A1'] = "Ingresos por Película (Top 10)"
    ws2['A1'].font = Font(bold=True, size=14)
    ws2.merge_cells('A1:D1')
    
    ws2['A3'] = "Película"
    ws2['B3'] = "Total Recaudación"
    ws2['C3'] = "Precio Full"
    ws2['D3'] = "Con Promoción"
    for cell in ['A3', 'B3', 'C3', 'D3']:
        ws2[cell].fill = header_fill
        ws2[cell].font = header_font
        ws2[cell].alignment = Alignment(horizontal='center')
    
    row = 4
    for item in fin_data.get('detalle', []):
        ws2[f'A{row}'] = item['titulo']
        ws2[f'B{row}'] = item['total']
        ws2[f'B{row}'].number_format = '$#,##0.00'
        ws2[f'C{row}'] = item['full']
        ws2[f'C{row}'].number_format = '$#,##0.00'
        ws2[f'D{row}'] = item['promo']
        ws2[f'D{row}'].number_format = '$#,##0.00'
        row += 1
    
    ws2.column_dimensions['A'].width = 50
    ws2.column_dimensions['B'].width = 20
    ws2.column_dimensions['C'].width = 18
    ws2.column_dimensions['D'].width = 18
    
    # Guardar a buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_pdf_financiero(kpis, fin_data, fecha_inicio, fecha_fin, chart_base64=None, requested_by=None, ranking_revpas=None, ranking_dia_revenue=None, promociones_data=None, filtros_resumen=None):
    """
    Genera PDF con reporte financiero ejecutivo.
    
    Args:
        kpis (dict): KPIs globales
        fin_data (dict): datos financieros por película
        fecha_inicio (str): fecha de inicio en formato ISO
        fecha_fin (str): fecha de fin en formato ISO
    
    Returns:
        BytesIO: buffer con el archivo PDF
    """
    # Obtener configuración del cine para el template
    from cine.models import ConfiguracionCine
    try:
        configuracion_cine = ConfiguracionCine.load()
    except Exception:
        configuracion_cine = None
    
    context = {
        'titulo': 'Reporte Financiero Ejecutivo',
        'fecha_inicio': fecha_inicio,
        'fecha_fin': fecha_fin,
        'kpis': kpis,
        'peliculas': fin_data.get('detalle', []),
        'promociones_data': promociones_data if promociones_data is not None else fin_data.get('detalle', []),
        'chart_base64': chart_base64,
        'requested_by': requested_by,
        'configuracion_cine': configuracion_cine,
        'ranking_revpas': ranking_revpas or [],
        'ranking_dia_revenue': ranking_dia_revenue or [],
        'filtros_resumen': filtros_resumen or 'Sin filtros adicionales',
    }
    
    html_string = render_to_string('reportes/pdf_financiero.html', context)
    
    buffer = BytesIO()
    pisa_status = pisa.CreatePDF(html_string, dest=buffer)
    
    if pisa_status.err:
        raise Exception('Error al generar PDF financiero')
    
    buffer.seek(0)
    return buffer


def generar_excel_operativo(occ_data, fecha_inicio, fecha_fin):
    """
    Genera archivo Excel con reporte operativo de salas.
    
    Args:
        occ_data (dict): datos de ocupación semanal
        fecha_inicio (str): fecha de inicio en formato ISO
        fecha_fin (str): fecha de fin en formato ISO
    
    Returns:
        BytesIO: buffer con el archivo Excel
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ocupación Semanal"
    
    # Header styling
    header_fill = PatternFill(start_color="10b981", end_color="10b981", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    
    ws['A1'] = "Reporte Operativo de Salas"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:D1')
    
    ws['A2'] = f"Período: {fecha_inicio} a {fecha_fin}"
    ws.merge_cells('A2:D2')
    
    ws['A4'] = "Día de la Semana"
    ws['B4'] = "Tickets Vendidos"
    ws['C4'] = "Capacidad Total"
    ws['D4'] = "% Ocupación"
    for cell in ['A4', 'B4', 'C4', 'D4']:
        ws[cell].fill = header_fill
        ws[cell].font = header_font
        ws[cell].alignment = Alignment(horizontal='center')
    
    row = 5
    for item in occ_data.get('detalle', []):
        ws[f'A{row}'] = item['dia']
        ws[f'B{row}'] = item['tickets']
        ws[f'B{row}'].number_format = '#,##0'
        ws[f'C{row}'] = item['capacidad']
        ws[f'C{row}'].number_format = '#,##0'
        ws[f'D{row}'] = item['porcentaje'] / 100
        ws[f'D{row}'].number_format = '0.00%'
        row += 1
    
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 15
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_pdf_operativo(occ_data, fecha_inicio, fecha_fin, chart_base64=None, heatmap_base64=None, marketing_base64=None, requested_by=None, dashboard_alertas=None, ranking_peliculas=None, ranking_horarios=None, ranking_salas=None, tiene_filtros_personalizados=False, filtros_resumen=None):
    """
    Genera PDF con reporte operativo de salas.
    
    Args:
        occ_data (dict): datos de ocupación semanal
        fecha_inicio (str): fecha de inicio en formato ISO
        fecha_fin (str): fecha de fin en formato ISO
        ranking_salas (list): ranking de salas por ocupación
        ranking_horarios (list): ranking de horarios por día de la semana
        tiene_filtros_personalizados (bool): indica si el usuario aplicó ordenamiento personalizado
    
    Returns:
        BytesIO: buffer con el archivo PDF
    """
    # Obtener configuración del cine para el template
    from cine.models import ConfiguracionCine
    from datetime import datetime
    
    try:
        configuracion_cine = ConfiguracionCine.load()
    except Exception:
        configuracion_cine = None
    
    # Convertir strings de fecha a objetos date para el template
    try:
        fecha_inicio_obj = datetime.strptime(fecha_inicio, '%Y-%m-%d').date() if isinstance(fecha_inicio, str) else fecha_inicio
    except Exception:
        fecha_inicio_obj = fecha_inicio
    
    try:
        fecha_fin_obj = datetime.strptime(fecha_fin, '%Y-%m-%d').date() if isinstance(fecha_fin, str) else fecha_fin
    except Exception:
        fecha_fin_obj = fecha_fin
    
    # Verificar si hay datos válidos para horarios y ocupación
    tiene_horarios_validos = False
    if ranking_horarios:
        for dia in ranking_horarios:
            if dia.get('mejor_horario') != 'N/A':
                tiene_horarios_validos = True
                break
    
    tiene_ocupacion_valida = False
    ocupacion_detalle = occ_data.get('detalle', [])
    if ocupacion_detalle:
        for dia in ocupacion_detalle:
            if dia.get('capacidad', 0) > 0:
                tiene_ocupacion_valida = True
                break
    
    context = {
        'titulo': 'Reporte Operativo de Salas',
        'fecha_inicio': fecha_inicio_obj,
        'fecha_fin': fecha_fin_obj,
        'ocupacion': ocupacion_detalle,
        'chart_base64': chart_base64,
        'heatmap_base64': heatmap_base64,
        'marketing_base64': marketing_base64,
        'requested_by': requested_by,
        'configuracion_cine': configuracion_cine,
        'dashboard_alertas': dashboard_alertas or {},
        'ranking_peliculas': ranking_peliculas or [],
        'ranking_horarios': ranking_horarios or [],
        'ranking_salas': ranking_salas or [],
        'tiene_horarios_validos': tiene_horarios_validos,
        'tiene_ocupacion_valida': tiene_ocupacion_valida,
        'tiene_filtros_personalizados': tiene_filtros_personalizados,
        'filtros_resumen': filtros_resumen or 'Sin filtros adicionales'
    }
    
    html_string = render_to_string('reportes/pdf_operativo.html', context)
    
    buffer = BytesIO()
    pisa_status = pisa.CreatePDF(html_string, dest=buffer)
    
    if pisa_status.err:
        raise Exception('Error al generar PDF operativo')
    
    buffer.seek(0)
    return buffer
